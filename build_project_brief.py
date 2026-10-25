from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Joaol\Documents\ChatGPT\ALIVE")
ASSETS = ROOT / "assets"
OUT = ROOT / "Trait_Graph_Coevolution_Project_Brief.docx"
HERO = ASSETS / "trait-graph-concept.png"
SCHEMATIC = ASSETS / "trait-graph-schematic.png"

NAVY = "102A43"
BLUE = "167D9A"
CYAN = "27B7D6"
AMBER = "E49A27"
INK = "243B53"
MUTED = "627D98"
PALE = "EDF6F8"
LIGHT = "F4F6F9"
WHITE = "FFFFFF"
RED = "A23B3B"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_font(run, name="Aptos", size=11, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (fld_char, instr, sep, text, end):
        run._r.append(node)
    set_font(run, size=9, color=MUTED)


def set_picture_alt(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        set_font(p.add_run(text))
    return p


def add_callout(doc, label, text, fill=PALE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label.upper() + "  ")
    set_font(r, size=9.5, bold=True, color=accent)
    r = p.add_run(text)
    set_font(r, size=11, bold=True, color=NAVY)
    set_table_geometry(table, [9360], indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def font(size, bold=False):
    try:
        return ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def create_schematic():
    canvas = Image.new("RGB", (1600, 650), "#F7FAFC")
    draw = ImageDraw.Draw(canvas)
    nodes = {
        "Foundation": (800, 540, NAVY),
        "Efficiency": (500, 390, CYAN),
        "Intensity": (1100, 390, AMBER),
        "Mobility": (280, 225, CYAN),
        "Sensing": (650, 225, CYAN),
        "Pressure": (950, 225, AMBER),
        "Contesting": (1320, 225, AMBER),
        "Persistence": (160, 80, BLUE),
        "Detection": (620, 80, BLUE),
        "Targeting": (980, 80, "8B5FBF"),
        "Resource control": (1435, 80, AMBER),
    }
    links = [
        ("Foundation", "Efficiency"), ("Foundation", "Intensity"),
        ("Efficiency", "Mobility"), ("Efficiency", "Sensing"),
        ("Intensity", "Pressure"), ("Intensity", "Contesting"),
        ("Mobility", "Persistence"), ("Sensing", "Detection"),
        ("Pressure", "Targeting"), ("Contesting", "Resource control"),
        ("Sensing", "Targeting"),
    ]
    for a, b in links:
        ax, ay, _ = nodes[a]
        bx, by, _ = nodes[b]
        draw.line((ax, ay, bx, by), fill="#9FB3C8", width=6)
    for label, (x, y, color) in nodes.items():
        radius = 48 if label == "Foundation" else 39
        draw.ellipse((x-radius, y-radius, x+radius, y+radius), fill="#" + color, outline="white", width=5)
        box = draw.textbbox((0, 0), label, font=font(25, bold=True))
        tx = x - (box[2] - box[0]) / 2
        ty = y + radius + 10
        draw.rounded_rectangle((tx-12, ty-5, tx+(box[2]-box[0])+12, ty+(box[3]-box[1])+8), radius=8, fill="white")
        draw.text((tx, ty), label, fill="#" + INK, font=font(25, bold=True))
    canvas.save(SCHEMATIC, quality=95)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    bullet = styles["List Bullet"]
    bullet.font.name = "Aptos"
    bullet.font.size = Pt(11)
    bullet.paragraph_format.left_indent = Inches(0.5)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(6)
    bullet.paragraph_format.line_spacing = 1.1


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("RESEARCH CONCEPT BRIEF")
    set_font(hr, size=8.5, bold=True, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Trait-Graph Co-Evolution   |   ")
    set_font(fr, size=9, color=MUTED)
    add_page_field(fp)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("MASTER'S PROJECT CONCEPT")
    set_font(r, size=9.5, bold=True, color=AMBER)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    r = title.add_run("Trait-Graph Co-Evolution")
    set_font(r, name="Aptos Display", size=28, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("Can reciprocal evolution keep producing new, useful specializations?")
    set_font(r, size=14, color=BLUE)
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_after = Pt(10)
    shape = pic_p.add_run().add_picture(str(HERO), width=Inches(6.5))
    set_picture_alt(shape, "Trait-graph co-evolution concept", "Two abstract populations compete on a resource grid while branching networks represent evolving trait pathways.")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    set_font(cap.add_run("Concept visualization: two populations explore branching evolutionary pathways under finite resources."), size=8.5, italic=True, color=MUTED)
    add_callout(doc, "Core question", "Does reciprocal evolution produce deeper, more diverse, and more persistent trait pathways than one-sided evolution?")
    add_heading(doc, "The idea in plain language", 1)
    add_body(doc, "Two groups of digital agents compete for limited resources. Agents inherit traits, mutations unlock nearby branches in a trait graph, and natural selection determines which branches survive. The study compares one-sided evolution with reciprocal evolution to test whether competition creates lasting innovation or simply repeats old strategies.")


def add_trait_graph_page(doc):
    doc.add_page_break()
    add_heading(doc, "How the trait graph works", 1)
    add_body(doc, "The graph acts like an evolutionary possibility map. It gives agents room to specialize without pretending they can invent unlimited traits from nothing.")
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_after = Pt(4)
    shape = pic.add_run().add_picture(str(SCHEMATIC), width=Inches(6.5))
    set_picture_alt(shape, "Example trait graph", "A foundation trait branches toward efficiency and intensity, then toward several specialized functional capabilities.")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    set_font(cap.add_run("Illustrative structure only; final nodes and trade-offs would be calibrated before confirmatory runs."), size=8.5, italic=True, color=MUTED)
    add_heading(doc, "Evolutionary rules", 2)
    add_bullet(doc, "Mutation proposes change. Agents may activate a neighboring node, modify its strength, duplicate a module, or lose an active node.", "Mutation proposes change.")
    add_bullet(doc, "Selection decides what persists. Useful combinations improve survival or reproduction; ineffective combinations disappear.", "Selection decides what persists.")
    add_bullet(doc, "Every capability has a cost. Active nodes consume energy so larger graphs are not automatically stronger.", "Every capability has a cost.")
    add_bullet(doc, "No universal best path. Benefits depend on resources, population density, and the opposing population's current traits.", "No universal best path.")
    add_heading(doc, "Why a graph instead of unlimited traits?", 2)
    add_body(doc, "A completely unrestricted trait generator is difficult to interpret: extra genome components may be useless clutter rather than innovation. A branching graph keeps trait effects functional, measurable, and computationally compatible with fixed-shape JAX arrays while still allowing multiple evolutionary pathways.")


def add_experiment_page(doc):
    doc.add_page_break()
    add_heading(doc, "Proposed experiment", 1)
    add_body(doc, "All conditions begin from the same starting populations, resource layout, policy checkpoints, and paired random seeds. Only the allowed evolutionary pathway changes.")
    headers = ["Condition", "Population A", "Population B", "Purpose"]
    rows = [
        ["Frozen", "No evolution", "No evolution", "Ecological baseline"],
        ["Neutral drift", "Neutral mutation", "Neutral mutation", "Mutation-only comparison"],
        ["A-only", "Evolves", "Frozen", "One-sided adaptation"],
        ["B-only", "Frozen", "Evolves", "Reverse one-sided control"],
        ["Reciprocal", "Evolves", "Evolves", "Primary treatment"],
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_font(p.add_run(text), size=9.5, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            if row_idx == 4:
                shade_cell(cells[idx], PALE)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            set_font(p.add_run(text), size=9.2, bold=(idx == 0 or row_idx == 4), color=NAVY if row_idx == 4 else INK)
    set_table_geometry(table, [1650, 1950, 1950, 3810], indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_heading(doc, "What would be measured", 2)
    add_bullet(doc, "Branch depth: how far populations move from their starting capabilities.", "Branch depth:")
    add_bullet(doc, "Functional diversity: how many meaningfully different trait combinations coexist.", "Functional diversity:")
    add_bullet(doc, "Persistence: whether new branches survive across multiple population turnovers.", "Persistence:")
    add_bullet(doc, "Historical novelty: whether later populations differ functionally from their ancestors.", "Historical novelty:")
    add_bullet(doc, "Cross-play: whether current populations outperform archived opponents or merely cycle back to old strategies.", "Cross-play:")
    add_callout(doc, "Primary comparison", "Reciprocal evolution versus the average of both one-sided conditions, evaluated across paired seeds and multiple pretrained policy checkpoints.", fill=LIGHT)


def add_value_page(doc):
    doc.add_page_break()
    add_heading(doc, "Scientific value and honest limits", 1)
    add_heading(doc, "Why this is worth studying", 2)
    add_bullet(doc, "It tests a clear causal question about reciprocal versus one-sided evolution.")
    add_bullet(doc, "The trait graph creates visible evolutionary pathways that can be explained and analyzed.")
    add_bullet(doc, "Cross-play can distinguish continued escalation from recurring evolutionary cycles.")
    add_bullet(doc, "The simulator can become a platform for later work on environmental modification and more flexible genomes.")
    add_heading(doc, "What the project cannot claim", 2)
    add_bullet(doc, "A finite, researcher-designed graph is not proof of unlimited open-ended evolution.")
    add_bullet(doc, "More active traits do not automatically mean greater complexity or intelligence.")
    add_bullet(doc, "Results may depend on graph topology, trait costs, and resource rules, so these must be fixed before confirmatory runs.")
    add_heading(doc, "Recommended scope", 2)
    add_body(doc, "Use approximately 20-30 trait nodes, including trade-offs, incompatible branches, and frequency-dependent benefits. Begin with pretrained trait-conditioned policies frozen during the primary experiment. Treat continuous online learning as a later extension so learning instability does not obscure the evolutionary result.")
    add_heading(doc, "Project phases", 2)
    phases = [
        ("1", "Ecology", "Build a stable fixed-shape environment and verify energy accounting."),
        ("2", "Behavior", "Train and validate trait-conditioned policies with evolution disabled."),
        ("3", "Evolution", "Add graph mutation, inheritance, costs, and turnover tracking."),
        ("4", "Validation", "Test metrics on synthetic saturation, drift, expansion, and cycling."),
        ("5", "Experiment", "Freeze the design, preregister outcomes, and run confirmatory seeds."),
    ]
    table = doc.add_table(rows=0, cols=3)
    for i, name, detail in phases:
        cells = table.add_row().cells
        shade_cell(cells[0], BLUE)
        shade_cell(cells[1], PALE)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[2].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p0.add_run(i), size=12, bold=True, color=WHITE)
        set_font(cells[1].paragraphs[0].add_run(name), size=10, bold=True, color=NAVY)
        set_font(cells[2].paragraphs[0].add_run(detail), size=9.5, color=INK)
    set_table_geometry(table, [600, 1700, 7060], indent_dxa=120)
    add_heading(doc, "Advisor discussion point", 2)
    add_callout(doc, "Decision", "Is a bounded trait graph the right balance between scientific ambition, computational feasibility, and a defensible Master's thesis?", fill="FFF5E6", accent=AMBER)


def build():
    ASSETS.mkdir(parents=True, exist_ok=True)
    create_schematic()
    doc = Document()
    setup_styles(doc)
    configure_section(doc.sections[0])
    add_cover(doc)
    add_trait_graph_page(doc)
    add_experiment_page(doc)
    add_value_page(doc)
    props = doc.core_properties
    props.title = "Trait-Graph Co-Evolution Project Brief"
    props.subject = "Master's research project concept"
    props.author = "Research Project Team"
    props.keywords = "co-evolution, trait graph, natural selection, simulation"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
