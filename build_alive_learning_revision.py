from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from build_changing_worlds_proposal import geometry, margins, shade, repeat_header, alt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'ALIVE_Learning_and_Evolution_Revised_Proposal.docx'
NAVY, BLUE, MUTED = '153451', '2E74B5', '607286'
# grant_proposal preset; source-inspired teal accent and centered title are named overrides.
# Source-compact override: body 11 pt / 1.15 with 5 pt after; lists 1.10.
# Table text: 9.5 pt / 1.08 spacing; captions: 9 pt; equation block: 10 pt.
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)
for name, size, col, before, after in [
    ('Normal',11,NAVY,0,6),('Title',28,NAVY,0,8),('Subtitle',14,'167D9A',0,8),
    ('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',12,NAVY,8,4),
    ('List Bullet',11,NAVY,0,4),('Caption',9,MUTED,4,4)]:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(col)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.line_spacing=1.25
    st.paragraph_format.widow_control=True
    if name.startswith('Heading'): st.font.bold=True; st.paragraph_format.keep_with_next=True
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
st=doc.styles['List Bullet']; st.paragraph_format.left_indent=Inches(.375); st.paragraph_format.first_line_indent=Inches(-.194); st.paragraph_format.line_spacing=1.10
numroot=doc.part.numbering_part.element
for ab in numroot.findall(qn('w:abstractNum')):
    for lvl in ab.findall(qn('w:lvl')):
        ps=lvl.find(qn('w:pStyle'))
        if ps is not None and ps.get(qn('w:val'))=='ListBullet':
            pp=lvl.find(qn('w:pPr'))
            if pp is None: pp=OxmlElement('w:pPr'); lvl.append(pp)
            ind=pp.find(qn('w:ind'))
            if ind is None: ind=OxmlElement('w:ind'); pp.append(ind)
            ind.set(qn('w:left'),'540'); ind.set(qn('w:hanging'),'279')
            tabs=pp.find(qn('w:tabs'))
            if tabs is None: tabs=OxmlElement('w:tabs'); pp.append(tabs)
            for ch in list(tabs): tabs.remove(ch)
            tab=OxmlElement('w:tab'); tab.set(qn('w:val'),'num'); tab.set(qn('w:pos'),'540'); tabs.append(tab)
for name,size in [('Table Text',9.5),('Equation',10),('Small',9.5)]:
    st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH); st.base_style=doc.styles['Normal']; st.font.size=Pt(size)
    st.paragraph_format.line_spacing=1.08 if name=='Table Text' else 1.15
    st.paragraph_format.space_after=Pt(4)
    if name=='Equation': st.font.name='Consolas'
hp=sec.header.paragraphs[0]; hp.text='ALIVE  |  LEARNING AND EVOLUTION'; hp.style='Small'; hp.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Revised proposal  |  September 2026  |  ').font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)

def p(text='',style=None): return doc.add_paragraph(text,style)
def h(text,level=1): return p(text,f'Heading {level}')
def b(text): return p(text,'List Bullet')
def page(title): h(title).paragraph_format.page_break_before=True
def note(label,text):
    pa=p(); pa.add_run(label+'  ').bold=True; pa.add_run(text)
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EEF5F8'); pa._p.get_or_add_pPr().append(sh)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
    for c,s in zip(t.rows[0].cells,headers):
        c.paragraphs[0].style='Table Text'; r=c.paragraphs[0].add_run(s); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); shade(c,'F4F6F9')
    repeat_header(t.rows[0])
    for row in rows:
        cells=t.add_row().cells
        for c,s in zip(cells,row): c.paragraphs[0].style='Table Text'; c.paragraphs[0].add_run(s)
    geometry(t,widths)
    for row in t.rows:
        pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        for c in row.cells: margins(c,top=80,bottom=80)
    return t

# 1. Cover / abstract: retain the team's title, attribution and ALIVE expansion.
pa=p('ALIVE','Title'); pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
pa=p('Adaptive Life in Virtual Environments','Subtitle'); pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
pa=p('Learning Within Lifetimes,\nEvolving Across Generations','Title'); pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in pa.runs: r.font.size=Pt(24)
pa=p('A controlled study of learning, inherited adaptation,\nand an expandable modular trait graph','Subtitle'); pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Two-semester master\'s project | Revised September 2, 2026','Small')
p('Original proposal prepared by Akshay Kumar. Revised for the ALIVE team following advisor-style review.','Small')
h('Abstract')
p('ALIVE will investigate how learning during an individual lifetime interacts with inherited adaptation across generations. Digital agents will compete for finite, regrowing resources while researchers control the frequency of environmental change. A four-condition comparison will enable or disable lifetime learning and cumulative inherited adaptation, with stable-world controls and paired random seeds. The initial system will use eight ability modules, four continuous traits, and a small learner that updates preferences for exploration, harvesting, contesting, and conservation. Primary outcomes will measure sustained population performance and inherited module diversity. Standardized tests of archived designs will separate inherited improvement from benefits that require lifetime experience. The study extends established work on learning and evolution by examining these mechanisms within explicit modular constraints; a matched prerequisite-removal comparison will test the contribution of those constraints. Development will begin with an early complete pilot and expand through validation gates toward a larger trait graph. The intended result is a reproducible account of learning-evolution interactions within the tested worlds and architectures, including null effects, trade-offs, and extinction.')
note('Central question','How does the speed of environmental change affect the interaction between lifetime learning and inherited adaptation in agents with modular abilities?')

page('1. Research contribution and questions')
p('Learning can guide evolution without learned behavior being directly inherited. Hinton and Nowlan demonstrated this interaction in 1987 [1]. Aguilar, Bennati, and Helbing later studied foraging agents in cyclically changing environments and showed that learning can alter which inherited strategies evolve [2]. ALIVE therefore treats the interaction itself as established research territory.')
p('The proposed contribution is a controlled extension focused on modular abilities: prerequisites, energy costs, and restricted combinations may change how learning affects performance and persistent inherited diversity. This is a candidate contribution to substantiate through a focused literature comparison before design lock. A larger simulator or a larger trait tree alone does not establish scientific novelty.')
h('Primary questions',2)
b('Does the effect of combining learning and inherited adaptation differ between slowly and rapidly changing resource worlds?')
b('Does learning preserve several persistent module configurations, or favor a narrower set of flexible inherited designs?')
h('Mechanism and architecture questions',2)
b('Do later inherited designs perform better from birth under standardized conditions, and how much of their performance still depends on learning?')
b('Does removing prerequisite edges, while preserving module functions and costs, change the measured interaction?')
p('Predictions are conditional. Learning may improve short-term response when information remains useful long enough to act on it. Costs, misleading feedback, or frequent resets may remove that advantage. Evolution may favor specialization during long stable periods or flexible designs under recurrent change. The combined condition is not assumed to win.')
h('Scope of inference',2)
p('The core experiment estimates effects for the implemented learner, graph, resource rules, and schedules. It does not establish biological realism, human behavior, unlimited trait invention, or a universal advantage of learning. The project retains the synthetic-agent framing; language-model and digital-twin studies are background inspiration rather than evidence for its research gap.')
note('Minimum scientific result','Estimate the learning-by-inherited-adaptation interaction with uncertainty, and test whether archived designs support the proposed explanation.')

page('2. Simulator and lifetime learning')
p('A fixed-capacity JAX grid simulation will track agent position, energy, age, inherited design, and lifetime state. Resources regrow under researcher-defined rules. Movement, harvesting, metabolic costs, automatic energy-gated reproduction, and death use explicit event accounting. Newborn memory resets, and learned preference values are never inherited in the core study.')
p('Two labeled populations remain for continuity with ALIVE. Both use the same rules and experimental setting; labels do not confer abilities or favoritism. They are reporting groups within one interacting system, not independent replicates or evidence of reciprocal coevolution. A one-population smoke test will check that labels do not change behavior.')
h('Small, explicit learner',2)
p('Each agent stores four values Q: explore, harvest, contest, and conserve. A common fixed controller translates a selected mode into actions using local observations, energy, and active modules. All conditions use the same controller and action-availability rules. Learning-disabled agents retain the initial Q0 values; they can still react to observations.')
p('At each mode decision, select an available mode using an epsilon mixture of uniform exploration and softmax(Q / temperature). Hold the mode for H steps unless death or an invalid action interrupts it. Accumulate reward over the executed interval, then apply:')
p('Q_decay = (1 - decay) * Q + decay * Q0\nQ_next[mode] = Q_decay[mode] + alpha * (reward - Q_decay[mode])','Equation')
p('Other entries retain Q_decay. Reward is the bounded, normalized net energy flow from resource acquisition and operating costs over the interval. A separate event ledger excludes reproductive transfers and avoids treating energy-cap clipping as failure. Death receives a fixed terminal penalty. Normalization, H, penalty, exploration, and all update parameters are fixed after pilot testing.')
p('A fixed learning maintenance cost is deducted once per simulation step for learning-enabled agents and included once in operating costs. Learning rate, decay, cost, temperature, and Q0 do not evolve in the core study. All agents reset Q at birth. A short mode horizon only partly addresses delayed benefits; conclusions remain specific to this learner.')
note('Required learner checks','Demonstrate successful harvesting, beneficial exploration with delayed payoff, switching after resource relocation, and correct reproduction/death accounting before population experiments. If these fail, revise and retest the learner during the pilot.')

page('3. Trait graph: small core, large-tree roadmap')
p('The first validated model will contain eight modules and four continuous traits: scale, efficiency, sensing, and storage. All agents retain basic movement and harvesting. Modules add or modify capabilities; continuous traits change their magnitude only where the relevant capability exists. Every benefit carries an explicit cost or trade-off.')
# Static scientific diagram, rendered programmatically for exact labels.
im=Image.new('RGB',(1600,660),'white'); dr=ImageDraw.Draw(im)
font=ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf',27)
bold=ImageFont.truetype(r'C:\Windows\Fonts\arialbd.ttf',29)
def box(x,y,w,hh,text,fill,fg='white'):
    dr.rounded_rectangle((x-w//2,y-hh//2,x+w//2,y+hh//2),radius=14,fill=fill)
    dr.multiline_text((x,y),text,font=bold if y<230 else font,fill=fg,anchor='mm',align='center',spacing=5)
box(800,65,360,75,'Basic agent','#153451')
branches=[(210,'Mobility','Efficient movement','Burst movement','#2484A2'),(600,'Sensing','Local detection','Directional sensing','#4676B2'),(1000,'Resource use','Efficient harvesting','Energy storage','#2D8C72'),(1390,'Interaction','Contesting','Selective contesting','#9C6B31')]
for x,title,first,second,col in branches:
    dr.line((800,103,x,168),fill='#A9BACA',width=5)
    dr.line((x,234,x,308),fill=col,width=5); dr.polygon([(x-8,296),(x+8,296),(x,311)],fill=col)
    dr.line((x,402,x,490),fill=col,width=5); dr.polygon([(x-8,479),(x+8,479),(x,494)],fill=col)
    box(x,202,300,66,title,col); box(x,355,330,94,first,'#EDF3F8','#153451'); box(x,540,330,94,second,'#EDF3F8','#153451')
graph=ROOT/'assets'/'alive_core_trait_tree.png'; graph.parent.mkdir(exist_ok=True); im.save(graph)
pa=p(); sh=pa.add_run().add_picture(str(graph),width=Inches(6.5)); alt(sh,'Eight-module starter graph','Basic agent connects to four family labels. Each family contains two modules linked by a prerequisite arrow. Family labels and root are not counted as modules.')
p('Illustrative version 1: four family labels, eight functional modules. Downward arrows specify prerequisites. Final effects and costs are validated during the pilot.','Caption')
table(['Stage','Target','Entry condition'],[
('Validated core','8 modules; 4 traits','Complete all four conditions and validate learner, inheritance, costs, and metrics.'),
('Expanded ALIVE','20-30 modules; up to 8 traits','Core analysis complete; add branches in small, separately validated batches.'),
('Large trait tree','40-60+ modules; deeper branches','Long-term target, contingent on evidence, throughput, and team capacity.')],[1900,2300,5160])
p('Expanded branches may include protection, recovery, and flexibility. Those continuous traits, plus intensity, restore the original eight-trait ambition. The graph can later support cross-branch prerequisites while remaining acyclic. These are planned research extensions, not guaranteed thesis deliverables.','Small')

page('4. Four conditions and inheritance controls')
table(['Condition','Lifetime update','Newborn inherited design'],[
('Fixed-reference agents','Off; Q = Q0','Drawn from fixed initial distribution P0'),
('Learning only','On; resets at birth','Drawn from the same fixed P0'),
('Inherited adaptation only','Off; Q = Q0','Copied from successful parent, with mutation'),
('Learning + inherited adaptation','On; resets at birth','Copied from successful parent, with mutation')],[2750,2300,4310])
h('What the inheritance switch means',2)
p('In all conditions, reproduction requires the same energy threshold and charges the same parental energy cost. Offspring receive the same starting energy and placement rule. In reference conditions, offspring genotype is sampled independently from P0 instead of copied from the parent. This disables cumulative parent-to-offspring genetic adaptation while preserving ecological births and deaths.')
p('Differential survival can still change the traits represented among living reference agents. The design therefore calls these fixed-reference birth distributions, not completely fixed living populations. Birth and survivor distributions are logged separately. Disabling mutation alone is not a no-adaptation control because selection can act on standing inherited variation.')
p('The inherited-adaptation effect includes inheritance, selection across generations, and mutation as a package. It is not interpreted as the isolated effect of mutation. A required small sensitivity study will retain parental inheritance but disable mutation in learning-on and learning-off runs to distinguish selection on initial variation from the addition of new variants.')
h('Mutation and accounting',2)
p('Mutation makes either a bounded continuous-trait change or one valid module addition/removal. Removing a prerequisite with active dependents is rejected; invalid proposals leave the genotype unchanged. Mutation rates, boundary handling, and rejection rates are recorded. Dead slots reset all lifetime state before reuse. Full-capacity reproduction uses one predetermined allocation rule, and saturation is logged.')
h('Fair comparisons',2)
p('Conditions share starting samples and precomputed external resource schedules. Separate random streams handle world forcing, action choice, births, and mutation so different event counts do not shift environmental randomness. Resource stocks may diverge through consumption; identical external rules do not imply identical realized worlds. Both labeled populations receive the same treatment.')

page('5. Environmental design and replication')
p('The primary worlds use two regimes: broad abundance and patchy scarcity. The same regime equations are used in stable and changing treatments. Moving resource opportunities remain an extension after the main result. Returning to the original regime is built into repeated cycles, allowing recovery to be measured repeatedly.')
table(['World treatment','Protocol','Purpose'],[
('Stable abundance','Constant broad resources','Control for elapsed time in an easy world'),
('Stable scarcity','Constant patchy resources','Control for elapsed time in a difficult world'),
('Slow change','Repeated abundance/scarcity cycles','Change slower than baseline turnover'),
('Fast change','The same regimes, shorter periods','Change faster than baseline turnover')],[2300,3660,3400])
p('Estimate baseline lifetime and turnover during calibration using the reference ecology. Choose slow and fast durations relative to that common benchmark; report realized lifetimes for every condition, but do not alter schedules separately by condition. Pilot targets are roughly 0.5-1 baseline lifetime per fast regime and 5-10 per slow regime, subject to feasibility.')
p('All treatments run for the same total number of steps. Choose durations so changing runs contain complete cycles and equal time in each regime. Balance starting regime across seed blocks. Agents do not observe a global regime label or future schedule. Lock durations and any transition smoothing before confirmatory runs.')
h('Run plan',2)
p('Cross four adaptation conditions with four world treatments. Target 24 seed blocks, paired across all 16 cells: 384 primary runs. A benchmark-supported range of 20-30 blocks gives 320-480 runs; finalize the number before confirmatory data collection using pilot variability and a stated precision target. Agents, time steps, and population labels are not independent replicates.')
h('Required focused comparisons',2)
b('Architecture: repeat all four adaptation conditions in the slow and fast worlds with prerequisite edges removed. Keep module effects, costs, mutation attempts, and initial common-valid designs matched. Interpret changed reachability as part of this intervention; report acceptance rates.')
b('Control sensitivity: in slow and fast worlds, test parental inheritance without mutation for learning on and off. Use the same preselected subset of 8-12 paired seeds for these focused comparisons; report their lower precision.')
p('With 8-12 blocks, architecture and control sensitivity add 64-96 and 32-48 runs respectively. Learning-cost sweeps and moving-resource worlds follow only if the core and required checks are complete.','Small')

page('6. Measurements and mechanism tests')
h('Two primary outcomes',2)
b('Population performance: mean living-agent count over a fixed horizon, divided by initial agent count. This measures sustained population presence, includes extinction, and uses a common denominator. Report pool saturation to detect ceiling effects.')
b('Inherited module diversity: time-averaged effective number of module configurations, exp(-sum(p * log(p))), where p is each configuration\'s frequency among living agents at a checkpoint. Extinction contributes zero; also report diversity conditional on survival and population size.')
p('Use module bitmasks for the primary diversity measure. Continuous-trait spread, persistent configurations, action diversity, reproduction, resource efficiency, and energy balance are secondary. Define persistence as frequency above a fixed threshold over a baseline-turnover window; freeze that threshold before main runs.')
h('Separate inheritance from experience',2)
p('Archive genotypes at predetermined checkpoints across all seed blocks. Evaluate early and late designs as newborn agents with equal age, energy, reset Q, and standardized resource worlds. Use equal sample sizes and a fixed reference competitor cohort. Disable reproduction and mutation during these short evaluations; record survival and resource acquisition.')
p('For each design, run learning enabled and disabled. Better late-design performance without learning supports inherited improvement. The difference between learning-on and learning-off performance measures dependence on learning. Smaller Q changes alone are insufficient evidence. Evaluate the same genotypes in both resource regimes and compare evolved designs from both learning histories.')
h('Recovery and statistical analysis',2)
p('Define adaptation lag relative to a preregistered fraction of the destination-regime stable-control performance, not the abundance performance before a scarcity shock. Recovery must persist for a fixed window. Runs that never recover are censored; extinction remains a reported outcome. Shock loss integrates the positive deficit below the same reference. Match elapsed time when comparing return cycles with stable controls.')
p('Estimate the interaction on the preregistered outcome scale: combined - learning-only - inherited-only + fixed-reference. Compute paired contrasts by seed block and report 95% confidence intervals using block resampling. A repeated-measures model with seed/run clustering can support period-level analysis. A negative interaction is not by itself proof of reduced selection; use archived-design tests and lineage evidence.')
p('Preregister primary contrasts across slow and fast schedules, multiplicity handling for the two outcomes, exclusion rules limited to verified implementation failures, and extinction handling. Report all seed outcomes. Sensitivity analyses supplement the registered definitions rather than selecting the most favorable one.','Small')

page('7. Two-semester implementation plan')
table(['Period','Work','Required milestone'],[
('Semester 1\nWeeks 1-2','Literature and minimal equations','Closest-work comparison, control definitions, resource and reward accounting'),
('Weeks 3-6','Small complete pilot','All four adaptation conditions; 8-module graph; learner and inheritance tests'),
('Weeks 7-10','Validate and benchmark','CPU/GPU throughput, stable ecology, memory reset, event ledger, synthetic metric traces'),
('Weeks 11-13','Calibrate study','Slow/fast regimes, reference-process sensitivity, parameter robustness, precision planning'),
('Weeks 14-16','Freeze version 1','Preregister graph, learner, endpoints, schedules, seeds, compute and storage budget'),
('Semester 2\nWeeks 1-5','Primary experiments','Complete paired matrix, checkpoints and preliminary registered analysis'),
('Weeks 6-8','Required mechanism checks','Archived-genotype tests, prerequisite removal, mutation-off sensitivity'),
('Weeks 9-11','Conditional expansion','Validate selected new branches if core evidence and budget permit'),
('Weeks 12-16','Analysis and writing','Final uncertainty estimates, limitations, thesis, reproducible release')],[1950,2600,4810])
h('Expansion gates',2)
b('Before 20-30 modules: version 1 runs reliably, learner tests pass, primary analysis is complete, and at least one explanation has been tested. A null result qualifies; a desired positive effect is not required.')
b('Add 2-4 modules at a time. Each batch needs a demonstrated function, a measurable cost, valid inheritance, and no unintended effect on old configurations. Rebenchmark and test on pilot seeds.')
b('Give each graph version immutable module IDs, effect definitions, prerequisite edges, cost tables, and compatibility tests. Store graph version in every result. Recompile fixed JAX shapes between versions when needed; do not change dimensions within a run.')
b('Large-tree experiments are a new registered study with their own controls and seeds. Compare conclusions across graph versions; do not merge their raw diversity counts or reuse the original version\'s claim without retesting.')

page('8. Compute, risks, and delivery')
h('Budget based on measured throughput',2)
p('The earlier 100-250 GPU-hour figure is retained only as a provisional planning envelope. It is not yet a measured requirement, and the larger trait tree has a separate budget. The core learner has no per-agent neural training loop; GPU need will depend on grid size, interaction rules, agent capacity, horizon, and batching.')
p('Benchmark the full simulator with learning, mutation, and realistic logging by semester 1, week 10. Report compilation time separately from warmed execution, along with hardware, grid size, active/capacity counts, memory use, and wall time per run. Compare CPU and GPU using identical workloads; small runs can remain on CPU.')
p('budget hours = run count * measured seconds per run / 3600','Equation')
p('The required simulation matrix contains approximately 416-624 runs before archived-design evaluations, calibration, and reruns. As arithmetic examples only, 10 minutes per run implies about 69-104 hours; 30 minutes implies 208-312 hours. Add separately measured evaluation and development costs plus a 20% contingency. Monetary cost equals final GPU-hours times the team\'s actual rate; no price is assumed.')
h('Main risks and responses',2)
b('Calibration bias: use pilot seeds to establish viable dynamics and functional mechanisms, not to make the combined condition win. Freeze the main design and use unseen confirmatory seeds.')
b('Learning or graph artifacts: retain the reward ledger, simple learner tests, matched prerequisite comparison, and archive evaluations. Interpret reduced learning, convergence, or collapse without changing the main world afterward.')
b('Unreliable timing or excessive scope: reduce graph expansion and optional sweeps first. Preserve the four-condition core and required controls; revise replication only before confirmatory runs and disclose the change.')
b('Storage bottlenecks: aggregate summaries on device, export at fixed checkpoints, and sample detailed agent traces. Save complete lineage events in batches where feasible and state any sampling limits.')
h('Deliverables and completion criteria',2)
p('Deliver the tested simulator, immutable version 1 graph, experiment configurations and seeds, throughput report, registered analysis, archived-design evaluations, uncertainty estimates, thesis, and reproducible figures. Deliver a documented expansion interface and large-tree roadmap even if the expanded graph is not built during the thesis. Scientific completion depends on valid estimates and honest interpretation, not a particular condition winning.')

page('9. Revision record and references')
h('Changes from the team proposal',2)
b('Preserved ALIVE, the team\'s learning/evolution question, modular abilities, and researcher-controlled environments. The long-term ambition remains a large trait tree.')
b('Replaced the initial 20-30-module commitment with a validated eight-module core and explicit expansion gates.')
b('Narrowed the contribution relative to existing learning-and-evolution research; added a direct prerequisite-removal comparison.')
b('Clarified the fixed-reference birth process, residual survival selection, and the combined meaning of the inheritance switch.')
b('Specified a candidate learner with reward accounting and validation tests; held learning parameters fixed initially.')
b('Made stable controls, slow/fast environmental change, archived-genotype evaluations, and paired run-level inference part of the required study.')
b('Moved the complete pilot to weeks 3-6 and replaced an unsupported compute estimate with benchmark-based budgeting.')
h('Selected research references',2)
p('[1] Hinton, G. E., & Nowlan, S. J. (1987). How Learning Can Guide Evolution. Complex Systems, 1, 495-502.','Small')
p('https://www.cs.toronto.edu/~hinton/absps/evolution.htm','Small')
p('[2] Aguilar, L., Bennati, S., & Helbing, D. (2019). How learning can change the course of evolution. PLOS ONE, 14(9), e0219502.','Small')
p('https://doi.org/10.1371/journal.pone.0219502','Small')
p('[3] Tang, Y., Tian, Y., & Ha, D. (2022). EvoJAX: Hardware-Accelerated Neuroevolution. arXiv:2202.05008. Engineering reference for accelerated evolutionary workloads; not a benchmark for this proposed simulator.','Small')
p('https://arxiv.org/abs/2202.05008','Small')
h('Design status',2)
p('This is a revised proposal, not a record of completed experiments. Module effects, numerical world settings, learner parameters, and final replication count will be fixed through the stated pilot and preregistration process. The proposed 40-60+ module target expresses the team\'s longer-term direction and is not a prerequisite for a successful thesis.','Small')
doc.core_properties.title='ALIVE: Learning Within Lifetimes, Evolving Across Generations - Revised Proposal'
doc.core_properties.subject='Validated core study and large trait-tree roadmap'
doc.core_properties.author='ALIVE team; original proposal by Akshay Kumar'
doc.save(OUT)
print(OUT)
